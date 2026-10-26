"""Step 4 request: what Hunter and Apollo should fetch for one company (hunter-apollo-v3/run_company.py --spec).

    python contact_spec.py --product tyres-pcr-2w-3w --country Peru --company "Comercio & Cía – Llantas"
        --names "COMERCIO & CIA S.A.;Grupo Molicom" --domains comerciollantas.com.pe;molicom.com.pe --out <batch>/spec.json

Everything comes from LeadGen_Memory (contact_filters.md for levels, departments and title/inbox words in the
country's language(s); the product profile's skip_inboxes) and from the workbook's Contacts tab:
  known_emails  every email already on the Contacts tab (never revealed or returned again)
  known_people  names already recorded on these domains (an Apollo candidate with that name is not revealed again)
Only own, verified domains go in --domains: the skill never searches by name alone in this mode."""
import argparse, json, os
from openpyxl import load_workbook
from leadgen_config import add_args, contact_filters, load


def known_contacts(cfg, domains):
    emails, people = set(), []
    if not os.path.exists(cfg.workbook):
        return emails, people
    try:
        ws = load_workbook(cfg.workbook, read_only=True)["Contacts"]
    except PermissionError:
        # Known contacts are what stop paid reveals of people already recorded: never build a request without them.
        raise SystemExit(f"{cfg.workbook} is open in Excel - close it and run contact_spec.py again "
                         "(the request needs the Contacts tab to skip people already recorded).")
    rows = ws.iter_rows(values_only=True)
    head = list(next(rows))
    ei, fi, li = head.index("Email"), head.index("First Name"), head.index("Last Name")
    for r in rows:
        email = (r[ei] or "").strip().lower()
        if not email:
            continue
        emails.add(email)
        if r[fi] and email.split("@")[-1] in domains:
            people.append([r[fi], r[li] or ""])
    return emails, people


def build(cfg, company, names, domains):
    f = contact_filters(cfg)
    s = f["settings"]
    lst = lambda k: [x.strip() for x in s.get(k, "").split(",") if x.strip()]
    domains = [d.lower().replace("https://", "").replace("http://", "").strip("/").removeprefix("www.") for d in domains]
    emails, people = known_contacts(cfg, set(domains))
    return {
        "company": company, "names": list(dict.fromkeys([company] + names)), "domains": domains,
        "country": cfg.country_name.replace("-", " "),
        "locations": [cfg.country_name.replace("-", " ")],  # Apollo: people based in the campaign country only
        "location_codes": [cfg.iso],  # Hunter's location filter takes ISO 3166-1 alpha-2 codes (country profile `iso`)
        "levels": lst("levels"),
        "hunter": {"purchasing_department": s.get("hunter_purchasing_department", "procurement"),
                   "seniority": lst("hunter_seniority"), "departments": lst("hunter_departments"),
                   "limit": int(s.get("hunter_limit", 10))},
        "budget": int(s.get("budget", 5)),
        "words": f["lists"],
        "known_emails": sorted(emails), "known_people": people,
    }


if __name__ == "__main__":
    p = add_args(argparse.ArgumentParser(description=__doc__.split("\n")[0]))
    p.add_argument("--company", required=True, help="trading name as on the Companies tab")
    p.add_argument("--names", default="", help="other names, ';'-separated: legal, group, brand, sister trade names")
    p.add_argument("--domains", required=True, help="own verified domains, ';'-separated")
    p.add_argument("--out", required=True)
    a = p.parse_args()
    spec = build(load(a.product, a.country), a.company, [n.strip() for n in a.names.split(";") if n.strip()],
                 [d.strip() for d in a.domains.split(";") if d.strip()])
    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    json.dump(spec, open(a.out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"{a.out}: {len(spec['names'])} names, {len(spec['domains'])} domains, levels {spec['levels']}, "
          f"{sum(map(len, spec['words'].values()))} filter words, {len(spec['known_emails'])} known emails, "
          f"{len(spec['known_people'])} known people on these domains")
