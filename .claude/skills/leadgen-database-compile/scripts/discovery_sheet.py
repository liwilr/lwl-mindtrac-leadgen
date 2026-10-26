"""Save a dealer / distributor list published by a researched company as an Excel sheet in the campaign folder's
additional_discovery/ subfolder: <Company>_<what the list contains>.xlsx. One row per listed dealer, as published.
No research happens here; each dealer goes through the normal workflow when it is picked from the list.

    python discovery_sheet.py --product tyres-pcr-2w-3w --country Peru --name "Maxxis Peru_sub-distributors" rows.json
    python discovery_sheet.py --product ... --country ... --refresh          re-check Status in every sheet against the registry

rows.json = [{"Name": ..., "City / Region": ..., "Address": ..., "Phone / WhatsApp": ..., "Email": ..., "Source URL": ...}, ...]
Status: "not researched" / "in registry: <class>" (set automatically) / "skipped: <reason>" (typed by hand, kept on refresh)."""
import _paths  # noqa: F401  (shared LeadGen_Tools on the path)
import argparse, csv, functools, glob, json, os, re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from leadgen_config import add_args, load, strip_accents

COLS = ["Name", "City / Region", "Address", "Phone / WhatsApp", "Email", "Source URL", "Date captured", "Status"]
WIDTHS = [38, 22, 46, 26, 34, 40, 14, 24]


def _norm(s):
    return re.sub(r"[^a-z0-9]+", " ", strip_accents((s or "").lower())).strip()


def _digits(s):
    return {re.sub(r"\D", "", m)[-9:] for m in re.findall(r"\d[\d \-]{7,}\d", s or "")}


@functools.lru_cache(None)
def _book(path):
    """Contacts tab as [(company, all cell text)] - read once per run."""
    if not os.path.exists(path):
        return []
    ws = load_workbook(path, read_only=True)["Contacts"]
    return [(r[0], " ".join(str(v) for v in r if v)) for r in ws.iter_rows(min_row=2, values_only=True)]


def status(cfg, row):
    if not os.path.exists(cfg.registry):
        return "not researched"
    book = _book(cfg.workbook)
    name, email, nums = _norm(row.get("Name")), (row.get("Email") or "").lower(), _digits(row.get("Phone / WhatsApp"))
    for r in csv.DictReader(open(cfg.registry, encoding="utf-8-sig")):
        legal = _norm(r["legal_name"])
        for suf in sorted((_norm(x) for x in cfg.legal_suffixes), key=len, reverse=True):
            if suf and legal.endswith(" " + suf):
                legal = legal[: -len(suf)].strip()
                break
        names = [_norm(r["trading_name"]), legal]
        if any(n and len(n) >= 6 and (n in name or name in n) for n in names if name):
            return f"in registry: {r['classification']}"
        blob = " ".join(t for c, t in book if c == r["trading_name"])
        if (email and email in blob.lower()) or (nums & _digits(blob)):
            return f"in registry: {r['classification']}"
    return "not researched"


def write(cfg, name, rows):
    os.makedirs(cfg.discovery_dir, exist_ok=True)
    path = os.path.join(cfg.discovery_dir, name + ".xlsx")
    wb = Workbook(); ws = wb.active; ws.title = "List"
    for c, h in enumerate(COLS, 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", start_color="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[ws.cell(1, c).column_letter].width = WIDTHS[c - 1]
    for n, row in enumerate(rows, 2):
        row = {**row, "Date captured": row.get("Date captured") or cfg.today}
        if not (row.get("Status") or "").startswith("skipped"):
            row["Status"] = status(cfg, row)
        for c, h in enumerate(COLS, 1):
            cell = ws.cell(n, c, row.get(h))
            cell.font = Font(name="Arial", size=10); cell.alignment = Alignment(vertical="top", wrap_text=h != "Status")
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    return path


def read(path):
    ws = load_workbook(path)["List"]
    head = [c.value for c in ws[1]]
    return [dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True) if any(r)]


if __name__ == "__main__":
    p = add_args(argparse.ArgumentParser(description=__doc__.split("\n")[0]))
    p.add_argument("rows", nargs="?"); p.add_argument("--name"); p.add_argument("--refresh", action="store_true")
    a = p.parse_args()
    cfg = load(a.product, a.country)
    if a.refresh:
        for f in glob.glob(os.path.join(cfg.discovery_dir, "*.xlsx")):
            rows = read(f); write(cfg, os.path.basename(f)[:-5], rows)
            print(os.path.basename(f), {s: sum(1 for r in read(f) if r["Status"] == s) for s in sorted({r["Status"] for r in read(f)})})
    else:
        rows = json.load(open(a.rows, encoding="utf-8"))
        path = write(cfg, a.name, rows)
        out = read(path)
        print(path, "|", len(out), "rows |", {s: sum(1 for r in out if r["Status"] == s) for s in sorted({r["Status"] for r in out})})
