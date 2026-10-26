"""The one command registry-compile runs. Every write to a campaign's two result files goes through here.

    python compile.py apply          --product P --country C record.json [record.json ...]
    python compile.py new-campaign   --product P --country C          create the campaign folder, workbook and registry
    python compile.py refresh-overview --product P --country C        rewrite the workbook's Overview qualification section
    python compile.py dealer-list    --product P --country C --name "<Company>_<what it lists>" rows.json
    python compile.py refresh-dealers --product P --country C         update Status in additional_discovery/ sheets
    python compile.py status         --product P --country C          queued / failed records, last backup

apply = validate every record -> back up the workbook and registry -> queue the records -> apply them (apply.py) ->
read the written rows back and compare -> print a JSON report. Nothing is written if validation fails; a record that
can't be applied goes to queue/<campaign>/failed/ with its reason. Backups: queue/<campaign>/backups/<time>/ (last 10)."""
import _paths  # noqa: F401  (shared LeadGen_Tools on the path)
import argparse, contextlib, csv, datetime, glob, json, os, re, shutil, subprocess, sys
from openpyxl import load_workbook
from leadgen_config import REGISTRY_FIELDS, add_args, load
import apply as queue

HERE = os.path.dirname(os.path.abspath(__file__))
NEW_KEYS = {"company", "contacts", "registry"}
UPDATE_KEYS = {"update", "company_append", "company_set", "contacts", "registry_append", "registry_set", "move_contacts",
               "renumber", "kw_keys"}
CONTACT_KEYS = {"Email", "Type", "First Name", "Last Name", "Position", "Department", "Phone", "WhatsApp",
                "Phone / WhatsApp", "Branch / Location", "LinkedIn", "Source"}


def validate(rec):
    """-> list of problems (empty = valid)."""
    p = []
    if not isinstance(rec, dict):
        return ["not a JSON object"]
    if rec.get("kw_new_rows"):
        p.append("kw_new_rows present: keyword rows are written by lead-gen before compiling")
    if "update" in rec:
        extra = set(rec) - UPDATE_KEYS
        if extra:
            p.append(f"unknown keys for an update record: {sorted(extra)}")
    else:
        missing = NEW_KEYS - set(rec)
        if missing:
            p.append(f"new-company record without {sorted(missing)}")
        reg, comp = rec.get("registry") or {}, rec.get("company") or {}
        for k in ("trading_name", "classification"):
            if not reg.get(k):
                p.append(f"registry.{k} is empty")
        if comp.get("Classification") != reg.get("classification"):
            p.append("company Classification and registry classification differ")
        if reg.get("classification") == "Reject" and rec.get("contacts"):
            p.append("a Reject carries contacts (lead-gen drops a Reject's contacts)")
        unknown = set(reg) - set(REGISTRY_FIELDS)
        if unknown:
            p.append(f"unknown registry fields: {sorted(unknown)}")
    rn = rec.get("renumber")
    if rn is not None and (not isinstance(rn, dict) or not rn or not all(re.fullmatch(r"\d{6,15}", re.sub(r"\D", "", str(x))) for kv in rn.items() for x in kv)):
        p.append("renumber must map stored numbers to corrected numbers (digits, 6-15 each)")
    for c in rec.get("contacts") or []:
        unknown = set(c) - CONTACT_KEYS
        if unknown:
            p.append(f"contact {c.get('Email') or c.get('Phone')}: unknown fields {sorted(unknown)}")
        if c.get("Date Contacted"):
            p.append("Date Contacted is filled only when a contact is actually contacted")
    return p


def backup(cfg):
    root = os.path.join(cfg.queue_dir, "backups")
    dest = os.path.join(root, datetime.datetime.now().strftime("%Y%m%d-%H%M%S"))
    os.makedirs(dest, exist_ok=True)
    for f in (cfg.workbook, cfg.registry):
        if os.path.exists(f):
            shutil.copy2(f, dest)
    for old in sorted(glob.glob(os.path.join(root, "*")))[:-10]:
        shutil.rmtree(old, ignore_errors=True)
    return dest


def read_back(cfg, rec):
    """Compare what the record asked for with what is now in the workbook and registry -> list of mismatches."""
    wb = load_workbook(cfg.workbook, read_only=True)
    comp = list(wb["Companies"].iter_rows(values_only=True))
    cont = list(wb["Contacts"].iter_rows(values_only=True))
    reg = list(csv.DictReader(open(cfg.registry, encoding="utf-8-sig")))
    ch, th = list(comp[0]), list(cont[0])
    name = rec.get("update") or rec["registry"]["trading_name"]
    first = lambda v: str(v or "").split("\n")[0]
    rows = [r for r in comp[1:] if first(r[0]).startswith(name) or name.startswith(first(r[0]))]
    regs = [r for r in reg if r["trading_name"] == name] or [r for r in reg if r["trading_name"].startswith(name)]
    bad = []
    if not rows:
        return [f"{name}: not on the Companies tab"]
    if not regs:
        bad.append(f"{name}: not in the registry")
    want = (rec.get("company_set") or {}).get("Classification") or (rec.get("company") or {}).get("Classification")
    if want and rows[-1][ch.index("Classification")] != want:
        bad.append(f"{name}: Classification is {rows[-1][ch.index('Classification')]}, expected {want}")
    if want and regs and regs[0]["classification"] != want:
        bad.append(f"{name}: registry classification is {regs[0]['classification']}, expected {want}")
    emails = {str(r[th.index('Email')] or '').lower() for r in cont[1:]}
    for c in (rec.get("contacts") or []):
        if c.get("Email") and c["Email"].lower() not in emails:
            bad.append(f"{name}: contact {c['Email']} not on the Contacts tab")
    for e in rec.get("move_contacts") or []:
        if e.lower() not in emails:
            bad.append(f"{name}: moved contact {e} missing")
    if rec.get("renumber"):
        mine = [r for r in cont[1:] if first(r[0]) and (name.startswith(first(r[0])) or first(r[0]).startswith(name))]
        nums = {re.sub(r"\D", "", l) for r in mine for col in ("Phone", "WhatsApp") for l in str(r[th.index(col)] or "").splitlines()}
        for old, new in rec["renumber"].items():
            o, n = re.sub(r"\D", "", old), re.sub(r"\D", "", new)
            if n not in nums or (o in nums and o != n):
                bad.append(f"{name}: renumber {o} -> {n} not visible on its Contacts rows")
    return bad


def cmd_apply(cfg, files):
    report = {"applied": [], "failed": [], "invalid": [], "still_queued": 0, "mismatches": [], "backup": None}
    recs = []
    for f in files:
        try:
            rec = json.load(open(f, encoding="utf-8"))
        except ValueError as e:
            report["invalid"].append({"file": f, "problems": [f"broken JSON: {e}"]})
            continue
        probs = validate(rec)
        (report["invalid"].append({"file": f, "problems": probs}) if probs else recs.append((f, rec)))
    if report["invalid"]:
        report["note"] = "nothing written: fix the invalid records and send the batch again"
        return report
    report["backup"] = backup(cfg)
    queued = {queue.save(cfg, f): rec for f, rec in recs}
    before_failed = set(queue.failed_records(cfg))
    with contextlib.redirect_stdout(sys.stderr):  # keep stdout for the JSON report
        queue.apply_all(cfg)
    for path, rec in queued.items():
        base = os.path.basename(path)
        if os.path.exists(os.path.join(cfg.queue_dir, "applied", base)):
            report["applied"].append(base)
            report["mismatches"] += read_back(cfg, rec)
        elif os.path.join(cfg.queue_dir, "failed", base) in set(queue.failed_records(cfg)) - before_failed:
            err = open(os.path.join(cfg.queue_dir, "failed", base + ".error.txt"), encoding="utf-8").read().strip()
            report["failed"].append({"record": base, "reason": err})
    report["still_queued"] = len(queue.pending(cfg))
    if report["still_queued"]:
        report["note"] = "workbook or registry open in another program: records stay queued; run apply again later"
    return report


def run_script(name, args):
    p = subprocess.run([sys.executable, os.path.join(HERE, name)] + args, capture_output=True, text=True,
                       encoding="utf-8", errors="replace", env=dict(os.environ, PYTHONIOENCODING="utf-8"))
    return {"ok": p.returncode == 0, "output": (p.stdout + p.stderr).strip()}


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("command", choices=["apply", "new-campaign", "refresh-overview", "dealer-list", "refresh-dealers", "status"])
    add_args(ap)
    ap.add_argument("--name", help="dealer-list sheet name: <Company>_<what it lists>")
    ap.add_argument("files", nargs="*")
    a = ap.parse_args()
    base = ["--product", a.product, "--country", a.country]
    cfg = load(a.product, a.country)
    if a.command == "apply":
        out = cmd_apply(cfg, a.files)
    elif a.command == "new-campaign":
        out = run_script("new_campaign.py", base)
    elif a.command == "refresh-overview":
        out = run_script("new_campaign.py", base + ["--refresh-overview"])
    elif a.command == "dealer-list":
        out = run_script("discovery_sheet.py", base + ["--name", a.name] + a.files)
    elif a.command == "refresh-dealers":
        out = run_script("discovery_sheet.py", base + ["--refresh"])
    else:
        b = sorted(glob.glob(os.path.join(cfg.queue_dir, "backups", "*")))
        out = {"queued": len(queue.pending(cfg)), "failed": [os.path.basename(f) for f in queue.failed_records(cfg)],
               "last_backup": b[-1] if b else None}
    print(json.dumps(out, ensure_ascii=False, indent=1))
    sys.exit(0 if (a.command != "apply" or not (out["invalid"] or out["failed"] or out["mismatches"])) else 1)
