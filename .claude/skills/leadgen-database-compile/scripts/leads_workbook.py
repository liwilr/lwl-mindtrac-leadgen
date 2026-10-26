"""Write one company into a campaign's lead workbook and registry. Used by apply.py; columns are found by header name.

add_company(cfg, company, contacts, registry, after=None)   new company (after = trading name of a related company)
update_company(cfg, name, company_append, contacts, registry_append)   same legal entity found again"""
import _paths  # noqa: F401  (shared LeadGen_Tools on the path)
import copy, csv, re
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from contact_numbers import to_columns
from leadgen_config import REGISTRY_FIELDS

CLASS_FILL = {"High": "C6EFCE", "Medium": "FFEB9C", "Low": "F8CBAD", "Reject": "D9D9D9"}
BAND = "EEF3F8"
ONE_LINE = {"Classification": 15, "Date Added": 13, "Date Contacted": 15}  # fixed width, never wrapped
NUMBER_COLS = ("Phone", "WhatsApp")  # digits only, stored as text, one number per line


def _headers(ws):
    return [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]


def _col(headers, key):
    """Header position (0-based). 'Keywords' matches 'Keywords (ES)'."""
    for i, h in enumerate(headers):
        if h == key or (h or "").split(" (")[0] == key:
            return i
    raise KeyError(f"column '{key}' not found in {headers}")


def _row(headers, values):
    row = [None] * len(headers)
    for k, v in values.items():
        row[_col(headers, k)] = v
    return row


def _read(ws):
    rows = [[ws.cell(r, c).value for c in range(1, ws.max_column + 1)] for r in range(2, ws.max_row + 1)]
    return [r for r in rows if any(v is not None for v in r)]


def _rewrite(ws, rows, band):
    headers = _headers(ws)
    kc = _col(headers, "Classification")
    tmpl = [copy.copy(ws.cell(2, c)._style) for c in range(1, len(headers) + 1)]
    ws.delete_rows(2, ws.max_row)
    shade, prev = False, None
    for n, row in enumerate(rows, 2):
        if row[0] != prev:
            shade, prev = not shade, row[0]
        for c in range(1, len(headers) + 1):
            cell = ws.cell(n, c, row[c - 1])
            cell._style = copy.copy(tmpl[c - 1])
            if band:
                cell.fill = PatternFill("solid", start_color=BAND) if shade else PatternFill(fill_type=None)
        k = ws.cell(n, kc + 1)
        k.fill = PatternFill("solid", start_color=CLASS_FILL[k.value])
        k.font = Font(name="Arial", size=10, bold=True)
    for name in NUMBER_COLS:
        if name in headers:
            c = headers.index(name) + 1
            for r in range(2, ws.max_row + 1):
                ws.cell(r, c).number_format = "@"
                ws.cell(r, c).alignment = Alignment(vertical="top", wrap_text=True)
    for name, width in ONE_LINE.items():
        if name in headers:
            c = headers.index(name) + 1
            ws.column_dimensions[get_column_letter(c)].width = width
            for r in range(1, ws.max_row + 1):
                a = ws.cell(r, c).alignment
                ws.cell(r, c).alignment = Alignment(horizontal=a.horizontal, vertical=a.vertical or "top", wrap_text=False)
    if ws.title == "Companies":  # auto-height: no stored height, so Excel fits each row to its wrapped text
        for r in range(2, ws.max_row + 1):
            ws.row_dimensions[r].height = None
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"


def _contact(cfg, c):
    """Record contact -> Contacts-tab values. A free-text "Phone / WhatsApp" value (older record style) is split into
    the Phone and WhatsApp columns; the names/notes attached to its numbers are appended to Source. "Verification" is
    no longer a column and is ignored. Date Contacted stays blank until the contact is actually contacted."""
    c = {k: v for k, v in c.items() if k != "Verification"}
    prefix, mobile = cfg.country.get("phone_prefix", ""), cfg.country.get("mobile_pattern")
    if "Phone / WhatsApp" in c:
        default = "wa" if "whatsapp only" in (c.get("Type") or "").lower() else "phone"
        phone, wa, notes = to_columns(c.pop("Phone / WhatsApp"), prefix, default, mobile, cfg.country.get("trunk_prefix", "0"))
        c.setdefault("Phone", phone); c.setdefault("WhatsApp", wa)
        # WhatsApp display names only aid research; they are not recorded (user decision 2026-10-01)
        wa_nums = set((wa or "").split("\n"))
        kept = []
        for n in notes:
            num, _, txt = n.partition(" ")
            parts = [p.strip() for p in txt.split(";") if p.strip() and not re.match(r"(?i)(wa name:|personal wa)", p.strip())]
            if num in wa_nums:
                parts = [p for p in parts if not re.search(r"[A-Za-z]", p) or p.lower().endswith("number")]
            if parts:
                kept.append(f"{num} {'; '.join(parts)}")
        notes = kept
        if notes:
            c["Source"] = "; ".join(x for x in [c.get("Source"), "numbers: " + "; ".join(notes)] if x)
    for key in NUMBER_COLS:  # values given directly: already full numbers, just keep the digits
        if c.get(key):
            lines = [re.sub(r"\D", "", x) for x in re.split(r"[\n/;,]+", str(c[key]))]
            c[key] = "\n".join(x for x in lines if x) or None
    return c


def _first_line(v):
    return (v or "").split("\n")[0]


def _find_company(crows, name):
    hits = [n for n, r in enumerate(crows) if _first_line(r[0]).startswith(name) or name.startswith(_first_line(r[0]))]
    assert hits, f"company '{name}' not on the Companies tab"
    return hits[-1]


def read_registry(cfg):
    return list(csv.DictReader(open(cfg.registry, encoding="utf-8-sig")))


def write_registry(cfg, rows):
    with open(cfg.registry, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=REGISTRY_FIELDS, lineterminator="\n")
        w.writeheader(); w.writerows(rows)


def _check_closed(cfg):
    """Fails fast (PermissionError) if Excel has the workbook or the registry open, so nothing is half-written."""
    for path in (cfg.workbook, cfg.registry):
        open(path, "r+b").close()


def _save(cfg, wb):
    ov = wb["Overview"]
    for r in range(1, ov.max_row + 1):
        if ov.cell(r, 1).value == "Last updated":
            ov.cell(r, 2).value = max(str(ov.cell(r, 2).value or ""), cfg.today)
    wb.calculation.fullCalcOnLoad = True
    wb.save(cfg.workbook)


def add_company(cfg, company, contacts, registry, after=None):
    _check_closed(cfg)
    name, klass = registry["trading_name"], registry["classification"]
    wb = load_workbook(cfg.workbook)
    cs, ct = wb["Companies"], wb["Contacts"]
    ch, th = _headers(cs), _headers(ct)
    crows = _read(cs)
    crows.insert(_find_company(crows, after) + 1 if after else len(crows), _row(ch, company))
    _rewrite(cs, crows, False)
    trows = _read(ct)
    j = len(trows)
    if after and any(r[0] == after for r in trows):
        j = max(n for n, r in enumerate(trows) if r[0] == after) + 1
    trows[j:j] = [_row(th, {**_contact(cfg, c), "Company": name, "Classification": klass}) for c in contacts]
    _rewrite(ct, trows, True)
    _save(cfg, wb)
    rows = read_registry(cfg)
    a = max((n for n, r in enumerate(rows) if r["trading_name"] == after), default=len(rows) - 1) + 1 if after else len(rows)
    rows.insert(a, {f: registry.get(f, "") for f in REGISTRY_FIELDS})
    write_registry(cfg, rows)


def update_company(cfg, name, company_append=None, contacts=None, registry_append=None, company_set=None, registry_set=None,
                   move_contacts=None, renumber=None):
    """company_set / registry_set replace values (e.g. a reclassification: Classification, Keywords); a new
    Classification is copied to the company's Contacts rows too. *_append adds to the existing text.
    move_contacts = emails whose Contacts rows move to this company from another; every other field (Date Contacted
    included) is kept, only Company and Classification change.
    renumber = {"<number as stored>": "<corrected number>"}: replaces that number in the Phone / WhatsApp cells of this
    company's Contacts rows (e.g. a local number stored without the country code); every other field is kept."""
    _check_closed(cfg)
    rows = read_registry(cfg)  # found first, so a name problem stops the update before anything is written
    hits = [r for r in rows if r["trading_name"] == name] or [r for r in rows if r["trading_name"].startswith(name)]
    assert len(hits) == 1, f"registry: '{name}' matched {len(hits)} rows"
    reg = hits[0]
    wb = load_workbook(cfg.workbook)
    cs, ct = wb["Companies"], wb["Contacts"]
    ch, th = _headers(cs), _headers(ct)
    crows = _read(cs)
    ci = _find_company(crows, name)
    row = crows[ci]
    for key, text in (company_append or {}).items():
        i = _col(ch, key)
        row[i] = (row[i] + "\n" + text) if row[i] else text
    for key, value in (company_set or {}).items():
        row[_col(ch, key)] = value
    _rewrite(cs, crows, False)
    trows = _read(ct)
    klass = row[_col(ch, "Classification")]
    label = next((r[0] for r in trows if name.startswith(r[0]) or r[0].startswith(name)), _first_line(row[0]))
    kc, ec = _col(th, "Classification"), _col(th, "Email")
    moving = {e.lower() for e in (move_contacts or [])}
    moved = [r for r in trows if (r[ec] or "").lower() in moving and r[0] != label]
    # An email already on this company counts as moved, so a move record can be applied again safely.
    found = {(r[ec] or "").lower() for r in trows if (r[ec] or "").lower() in moving}
    assert found == moving, f"move_contacts: {sorted(moving - found)} not on the Contacts tab"
    trows = [r for r in trows if not any(r is m for m in moved)]
    for r in moved:
        r[0] = label
    mine = [n for n, r in enumerate(trows) if r[0] == label]
    if mine:
        j = mine[-1] + 1
    else:  # no contact rows yet: after the contacts of the companies above it on the Companies tab
        above = {_first_line(r[0]) for r in crows[:ci]}
        j = max((n for n, r in enumerate(trows) if any(a.startswith(r[0]) for a in above)), default=len(trows) - 1) + 1
    trows[j:j] = moved + [_row(th, {**_contact(cfg, c), "Company": label, "Classification": klass}) for c in (contacts or [])]
    for r in trows:
        if r[0] == label:
            r[kc] = klass
    if renumber:
        dig = lambda s: re.sub(r"\D", "", str(s or ""))
        old = {dig(k): dig(v) for k, v in renumber.items()}
        seen = set()
        for r in trows:
            if r[0] != label:
                continue
            for col in ("Phone", "WhatsApp"):
                i = _col(th, col)
                lines = str(r[i] or "").splitlines()
                if any(dig(l) in old for l in lines):
                    seen |= {dig(l) for l in lines if dig(l) in old}
                    r[i] = "\n".join(old.get(dig(l), l) if dig(l) in old else l for l in lines)
        assert seen == set(old), f"renumber: {sorted(set(old) - seen)} not on {label}'s Contacts rows"
    _rewrite(ct, trows, True)
    _save(cfg, wb)
    for k, v in (registry_append or {}).items():
        reg[k] = (reg[k] + ";" + v) if reg[k] else v
    for k, v in (registry_set or {}).items():
        reg[k] = v
    write_registry(cfg, rows)
