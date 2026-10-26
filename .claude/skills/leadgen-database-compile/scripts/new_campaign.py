"""Create a campaign folder for a product x country (Step A.5 of the method):
Database/<Product folder>/Country/<Country>/ with <Country>_<Product>_Leads.xlsx (Overview / Companies / Contacts),
a header-only <Country>_company_registry.csv and an empty additional_discovery/ subfolder. Existing files are left alone.

    python new_campaign.py --product tyres-pcr-2w-3w --country Brazil
    python new_campaign.py --product … --country … --refresh-overview   rewrite the Overview's qualification section
                                                                          of an existing workbook (Excel must be closed)"""
import _paths  # noqa: F401  (shared LeadGen_Tools on the path)
import argparse, csv, os, re
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from leadgen_config import CLASSES, COMPANY_HEADERS, CONTACT_HEADERS, REGISTRY_FIELDS, add_args, load
from leads_workbook import CLASS_FILL

BLUE, GREY = "1F4E78", "595959"
THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
HEAD = dict(font=Font(name="Arial", size=11, bold=True, color="FFFFFF"), fill=PatternFill("solid", start_color=BLUE),
            alignment=Alignment(horizontal="center", vertical="center", wrap_text=True), border=THIN)
BODY = dict(font=Font(name="Arial", size=10), alignment=Alignment(vertical="top", wrap_text=True), border=THIN)
FRAMEWORK = [  # (category, group prefix, Match meaning, Relevant meaning) - lead-gen skill, references/qualification.md
    ("Business role", "role", "Buys in bulk (importer, distributor, wholesaler, representative)", "Sells to or serves end users (retail, sales, servicing – B2C)"),
    ("Product", "product", "The product itself, or a brand of it", "Other products of the same industry, or their brands"),
    ("Segment", "segment", "Uses the product directly", "Uses a different range of the product"),
]
RULES = [
    ("High", "Role + Product + Segment, all Match (a multi-trade company is capped at Medium)"),
    ("Medium", "Two Match + the third Relevant, or two Match + the third not found"),
    ("Low", "One Match + the other two at least Relevant, or one Match + one Relevant + the third not found"),
    ("Reject", "Everything else (no Match, or one Match with nothing else)"),
]
CRITERIA_ROW = 16  # the qualification section starts here; rows above (title, info, summary) are never touched
KEYWORDS_NOTE = ("Matched terms by group (kw_match.py), e.g. Role: importer | Product: battery | Segment: car | Relevant: tyre | "
                 "Brands: VARTA. Rows added before 2026-10-06 hold the company's own wording instead.")


def _style(cell, **kw):
    for k, v in kw.items():
        setattr(cell, k, v)


def _sheet(wb, title, headers, widths, comments):
    ws = wb.create_sheet(title)
    for c, h in enumerate(headers, 1):
        _style(ws.cell(1, c, h), **HEAD)
        _style(ws.cell(2, c), **BODY)  # styled blank row: the write script copies its style
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
        if h in comments:
            ws.cell(1, c).comment = Comment(comments[h], "LeadGen")
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def write_criteria(ov, cfg):
    """Overview qualification section: the role / product / segment framework with example keywords, the class rules, notes."""
    kw = cfg.keywords()
    ex = lambda g: " / ".join(re.sub(r"\s*\([^)]*\)", "", vs[0]) for _, vs in kw[g][:5]) or "—"
    r = CRITERIA_ROW
    ov.cell(r, 1, "Qualification framework").font = Font(name="Arial", size=12, bold=True)
    for c, h in enumerate(("Category", "Match", "Relevant"), 1):
        _style(ov.cell(r + 1, c, h), **HEAD)
    for i, (name, key, m, rel) in enumerate(FRAMEWORK, r + 2):
        _style(ov.cell(i, 1, name), font=Font(name="Arial", size=10, bold=True), alignment=Alignment(vertical="top", wrap_text=True), border=THIN)
        _style(ov.cell(i, 2, f"{m}\ne.g. {ex(key + '_match')}"), **BODY)
        _style(ov.cell(i, 3, f"{rel}\ne.g. {ex(key + '_rel')}"), **BODY)
    r += 6
    ov.cell(r, 1, "Classification").font = Font(name="Arial", size=12, bold=True)
    for c, h in enumerate(("Class", "Rule"), 1):
        _style(ov.cell(r + 1, c, h), **HEAD)
    for i, (k, rule) in enumerate(RULES, r + 2):
        _style(ov.cell(i, 1, k), font=Font(name="Arial", size=10, bold=True), fill=PatternFill("solid", start_color=CLASS_FILL[k]),
               alignment=Alignment(vertical="top", wrap_text=True), border=THIN)
        _style(ov.cell(i, 2, rule), **BODY)
        ov.merge_cells(start_row=i, start_column=2, end_row=i, end_column=3)
    r += 7
    notes = [f"Brands count as the group set in the keyword file (a brand of the product = Product – Match). Legal suffixes do not "
             f"affect classification. Full keyword list: LeadGen_Memory/products/{cfg.product_key}/{os.path.basename(cfg.keyword_file)}",
             "The class is calculated from the keywords found (kw_match.py). It is overridden only for what keywords cannot show "
             "(Reject: a foreign company, or a maker of the product / its local subsidiary that sells only its own brand; a multi-trade company - "
             "its own pages list unrelated business lines - is capped at Medium); "
             "the reason is in the registry notes.",
             "Group and Relationship (Companies tab): Group holds only the group name (for filtering); Relationship holds the role and evidence. Connected companies (verified by >=2 independent signals) are kept as "
             "separate entities, classified on their own keywords, and placed next to each other."]
    for i, t in enumerate(notes, r):
        ov.cell(i, 1, t).font = Font(name="Arial", size=9, color=GREY)


def refresh_overview(cfg):
    """Rewrite the qualification section of an existing workbook's Overview (rows from CRITERIA_ROW down)."""
    open(cfg.workbook, "r+b").close()  # PermissionError if Excel has it open
    wb = load_workbook(cfg.workbook)
    ov = wb["Overview"]
    for m in [m for m in ov.merged_cells.ranges if m.min_row >= CRITERIA_ROW]:
        ov.unmerge_cells(str(m))
    for row in ov.iter_rows(min_row=CRITERIA_ROW, max_row=max(ov.max_row, CRITERIA_ROW + 20), max_col=6):
        for c in row:
            c.value = None
            c.style = "Normal"
    write_criteria(ov, cfg)
    for r in range(1, CRITERIA_ROW):
        if ov.cell(r, 1).value == "Last updated":
            ov.cell(r, 2).value = cfg.today
    for c in wb["Companies"][1]:
        if str(c.value or "").startswith("Keywords"):
            c.comment = Comment(KEYWORDS_NOTE, "LeadGen")
    wb.save(cfg.workbook)


def build_workbook(cfg):
    wb = Workbook()
    ov = wb.active; ov.title = "Overview"
    for col, w in zip("ABC", (16, 62, 60)):
        ov.column_dimensions[col].width = w
    ov["A1"] = f"Lead Generation – {cfg.country_name}"; ov["A1"].font = Font(name="Arial", size=16, bold=True, color=BLUE)
    info = [("Country", f"{cfg.country_name} (language: {cfg.lang}; company ID: {cfg.country.get('id_name') or 'n/a'})"),
            ("Ideal segment", cfg.product.get("ideal_segment") or ""), ("Last updated", cfg.today)]
    for r, (k, v) in enumerate(info, 3):
        ov.cell(r, 1, k).font = Font(name="Arial", size=10, bold=True)
        _style(ov.cell(r, 2, v), font=Font(name="Arial", size=10), alignment=Alignment(vertical="top", wrap_text=True))
        ov.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ov["A7"] = "Summary"; ov["A7"].font = Font(name="Arial", size=12, bold=True)
    for c, h in enumerate(("Classification", "Companies"), 1):
        _style(ov.cell(8, c, h), **HEAD)
    for r, k in enumerate(CLASSES, 9):
        _style(ov.cell(r, 1, k), font=Font(name="Arial", size=10, bold=True), fill=PatternFill("solid", start_color=CLASS_FILL[k]), border=THIN)
        _style(ov.cell(r, 2, f'=COUNTIF(Companies!$B$2:$B$5000,"{k}")'), font=Font(name="Arial", size=10), border=THIN)
    _style(ov.cell(13, 1, "Total researched"), font=Font(name="Arial", size=10, bold=True), border=THIN)
    _style(ov.cell(13, 2, "=SUM(B9:B12)"), font=Font(name="Arial", size=10, bold=True), border=THIN)
    note = Font(name="Arial", size=9, color=GREY)
    ov["A14"] = (f"Counts come from the Classification column on the Companies tab. Every researched company, including Rejects, "
                 f"is also logged in {os.path.basename(cfg.registry)}. Emails and phone / WhatsApp numbers are on the Contacts tab only.")
    ov["A14"].font = note
    write_criteria(ov, cfg)

    ch = [h if h != "Keywords" else f"Keywords ({cfg.lang.upper()})" for h in COMPANY_HEADERS]
    _sheet(wb, "Companies", ch, (30, 15, 22, 42, 44, 24, 16, 60, 13), {
        "Classification": "High / Medium / Low / Reject – see Overview tab for rules.",
        ch[COMPANY_HEADERS.index("Keywords")]: KEYWORDS_NOTE,
        "Group": "The group name only (e.g. El Sol group, Grupo Molicom), for filtering. Blank for standalone companies and for non-group links.",
        "Relationship": "Role in the group, parent/sister companies and the evidence (>=2 independent signals); or a non-group link (dealer of, name-related only). Connected companies are placed next to each other."})
    _sheet(wb, "Contacts", CONTACT_HEADERS, (18, 15, 36, 16, 12, 13, 22, 14, 16, 16, 34, 15, 38, 30), {
        "Phone": "Digits only with the country code (no +, spaces or brackets); one number per line.",
        "WhatsApp": "Digits only with the country code; confirmed WhatsApp numbers (wa.me business name or WA button). Unconfirmed numbers go in Phone.",
        "Date Contacted": "Date this contact was last contacted (YYYY-MM-DD); blank until contacted.",
        "Type": "Generic = role/branch inbox. Personal = named person; fill First/Last/Position/Department. Rows are grouped by company: HQ first, then named people, then branches.",
        "LinkedIn": "Personal contact's LinkedIn profile URL (from Apollo/Hunter or manual research).",
        "Source": "Where found: Website / Facebook / Instagram / Wayback (archive date) / Hunter / Apollo / Directory"})
    wb.save(cfg.workbook)


if __name__ == "__main__":
    p = add_args(argparse.ArgumentParser(description=__doc__.split("\n")[0]))
    p.add_argument("--refresh-overview", action="store_true", help="rewrite the qualification section of an existing workbook")
    a = p.parse_args()
    cfg = load(a.product, a.country)
    if a.refresh_overview:
        refresh_overview(cfg)
        print("Overview qualification section rewritten:", cfg.workbook)
        raise SystemExit
    os.makedirs(cfg.discovery_dir, exist_ok=True)
    for path, make in ((cfg.workbook, lambda: build_workbook(cfg)),
                       (cfg.registry, lambda: csv.writer(open(cfg.registry, "w", encoding="utf-8-sig", newline="")).writerow(REGISTRY_FIELDS))):
        if os.path.exists(path):
            print("exists, left alone:", path)
        else:
            make(); print("created:", path)
    print("folder:", cfg.campaign_dir)
